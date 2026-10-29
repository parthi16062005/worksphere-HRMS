import csv
import io
from decimal import Decimal

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy import case, func
from sqlalchemy.orm import Session

from auth import require_roles
from database import get_db
from models import User, Payroll, Employee

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


router = APIRouter(tags=["Payroll Reports"])


# =====================================================
# VALIDATE REPORT FILTERS
# =====================================================

def validate_report_filters(
    payroll_month: int | None,
    payroll_year: int | None,
):
    if payroll_month is not None:
        if payroll_month < 1 or payroll_month > 12:
            raise HTTPException(
                status_code=400,
                detail="Month must be between 1 and 12.",
            )

    if payroll_year is not None:
        if payroll_year < 2000 or payroll_year > 2100:
            raise HTTPException(
                status_code=400,
                detail="Year must be between 2000 and 2100.",
            )


# =====================================================
# PAYROLL SUMMARY REPORT
# HR / ADMIN ONLY
# =====================================================

@router.get("/payroll/reports/summary")
def get_payroll_summary(
    payroll_month: int | None = None,
    payroll_year: int | None = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_roles("HR", "ADMIN")
    ),
):
    validate_report_filters(payroll_month, payroll_year)

    query = db.query(Payroll)

    if payroll_month is not None:
        query = query.filter(
            Payroll.month == payroll_month
        )

    if payroll_year is not None:
        query = query.filter(
            Payroll.year == payroll_year
        )

    payroll_records = query.all()

    total_gross_salary = sum(
        (
            Decimal(str(record.gross_salary or 0))
            for record in payroll_records
        ),
        Decimal("0.00"),
    )

    total_deductions = sum(
        (
            Decimal(str(record.total_deductions or 0))
            for record in payroll_records
        ),
        Decimal("0.00"),
    )

    total_net_salary = sum(
        (
            Decimal(str(record.net_salary or 0))
            for record in payroll_records
        ),
        Decimal("0.00"),
    )

    paid_records = [
        record
        for record in payroll_records
        if record.status == "PAID"
    ]

    pending_records = [
        record
        for record in payroll_records
        if record.status == "PENDING"
    ]

    total_paid_salary = sum(
        (
            Decimal(str(record.net_salary or 0))
            for record in paid_records
        ),
        Decimal("0.00"),
    )

    total_pending_salary = sum(
        (
            Decimal(str(record.net_salary or 0))
            for record in pending_records
        ),
        Decimal("0.00"),
    )

    return {
        "total_payroll_records": len(payroll_records),
        "paid_payroll_count": len(paid_records),
        "pending_payroll_count": len(pending_records),
        "total_gross_salary": float(total_gross_salary),
        "total_deductions": float(total_deductions),
        "total_net_salary": float(total_net_salary),
        "total_paid_salary": float(total_paid_salary),
        "total_pending_salary": float(total_pending_salary),
    }


# =====================================================
# DEPARTMENT-WISE PAYROLL ANALYTICS
# HR / ADMIN ONLY
# =====================================================

@router.get("/payroll/reports/department")
def get_department_payroll_report(
    payroll_month: int | None = None,
    payroll_year: int | None = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_roles("HR", "ADMIN")
    ),
):
    validate_report_filters(payroll_month, payroll_year)

    query = (
        db.query(
            Employee.department.label("department"),

            func.count(
                Payroll.id
            ).label("payroll_count"),

            func.coalesce(
                func.sum(Payroll.gross_salary),
                0,
            ).label("total_gross_salary"),

            func.coalesce(
                func.sum(Payroll.total_deductions),
                0,
            ).label("total_deductions"),

            func.coalesce(
                func.sum(Payroll.net_salary),
                0,
            ).label("total_net_salary"),

            func.coalesce(
                func.sum(
                    case(
                        (
                            Payroll.status == "PAID",
                            Payroll.net_salary,
                        ),
                        else_=0,
                    )
                ),
                0,
            ).label("total_paid_salary"),

            func.coalesce(
                func.sum(
                    case(
                        (
                            Payroll.status == "PENDING",
                            Payroll.net_salary,
                        ),
                        else_=0,
                    )
                ),
                0,
            ).label("total_pending_salary"),
        )
        .join(
            Employee,
            Employee.id == Payroll.employee_id,
        )
    )

    if payroll_month is not None:
        query = query.filter(
            Payroll.month == payroll_month
        )

    if payroll_year is not None:
        query = query.filter(
            Payroll.year == payroll_year
        )

    results = (
        query
        .group_by(Employee.department)
        .order_by(Employee.department.asc())
        .all()
    )

    return [
        {
            "department": row.department,
            "payroll_count": row.payroll_count,
            "total_gross_salary": float(
                row.total_gross_salary
            ),
            "total_deductions": float(
                row.total_deductions
            ),
            "total_net_salary": float(
                row.total_net_salary
            ),
            "total_paid_salary": float(
                row.total_paid_salary
            ),
            "total_pending_salary": float(
                row.total_pending_salary
            ),
        }
        for row in results
    ]

# =====================================================
# EXPORT PAYROLL REPORT TO CSV
# HR / ADMIN ONLY
# =====================================================

@router.get("/payroll/reports/export/csv")
def export_payroll_csv(
    payroll_month: int | None = None,
    payroll_year: int | None = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_roles("HR", "ADMIN")
    ),
):
    # Validate the month and year.
    validate_report_filters(
        payroll_month,
        payroll_year,
    )

    # Join payroll records with employee details.
    query = (
        db.query(Payroll, Employee)
        .join(
            Employee,
            Employee.id == Payroll.employee_id,
        )
    )

    # Apply optional month filter.
    if payroll_month is not None:
        query = query.filter(
            Payroll.month == payroll_month
        )

    # Apply optional year filter.
    if payroll_year is not None:
        query = query.filter(
            Payroll.year == payroll_year
        )

    # Get the filtered records.
    records = (
        query
        .order_by(
            Payroll.year.desc(),
            Payroll.month.desc(),
            Employee.department.asc(),
            Employee.full_name.asc(),
        )
        .all()
    )

    # Create the CSV file in memory.
    output = io.StringIO(newline="")
    writer = csv.writer(output)

    # CSV column headings.
    writer.writerow([
        "Payroll ID",
        "Employee Code",
        "Employee Name",
        "Department",
        "Designation",
        "Month",
        "Year",
        "Basic Salary",
        "HRA",
        "DA",
        "Other Allowances",
        "Gross Salary",
        "PF Deduction",
        "Tax Deduction",
        "Other Deductions",
        "Total Deductions",
        "Net Salary",
        "Payment Status",
        "Payment Date",
    ])

    # Write payroll details.
    for payroll, employee in records:
        writer.writerow([
            payroll.id,
            employee.employee_code,
            employee.full_name,
            employee.department,
            employee.designation,
            payroll.month,
            payroll.year,
            payroll.basic_salary,
            payroll.hra,
            payroll.da,
            payroll.other_allowances,
            payroll.gross_salary,
            payroll.pf_deduction,
            payroll.tax_deduction,
            payroll.other_deductions,
            payroll.total_deductions,
            payroll.net_salary,
            payroll.status,
            (
                payroll.payment_date.isoformat()
                if payroll.payment_date
                else ""
            ),
        ])

    # Prepare the CSV response.
    csv_content = output.getvalue()
    output.close()

    filename = "payroll_report"

    if payroll_month is not None:
        filename += f"_month_{payroll_month:02d}"

    if payroll_year is not None:
        filename += f"_{payroll_year}"

    filename += ".csv"

    # Return a downloadable CSV file.
    return StreamingResponse(
        iter([csv_content]),
        media_type="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": (
                f'attachment; filename="{filename}"'
            )
        },
    )

# =====================================================
# EXPORT PAYROLL REPORT TO EXCEL
# HR / ADMIN ONLY
# =====================================================

@router.get("/payroll/reports/export/excel")
def export_payroll_excel(
    payroll_month: int | None = None,
    payroll_year: int | None = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(
        require_roles("HR", "ADMIN")
    ),
):
    # Validate filters.
    validate_report_filters(
        payroll_month,
        payroll_year,
    )

    # Query payroll records with employee details.
    query = (
        db.query(Payroll, Employee)
        .join(
            Employee,
            Employee.id == Payroll.employee_id,
        )
    )

    if payroll_month is not None:
        query = query.filter(
            Payroll.month == payroll_month
        )

    if payroll_year is not None:
        query = query.filter(
            Payroll.year == payroll_year
        )

    records = (
        query
        .order_by(
            Payroll.year.desc(),
            Payroll.month.desc(),
            Employee.department.asc(),
            Employee.full_name.asc(),
        )
        .all()
    )

    # Create an Excel workbook.
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Payroll Report"

    # Column headings.
    headers = [
        "Payroll ID",
        "Employee Code",
        "Employee Name",
        "Department",
        "Designation",
        "Month",
        "Year",
        "Basic Salary",
        "HRA",
        "DA",
        "Other Allowances",
        "Gross Salary",
        "PF Deduction",
        "Tax Deduction",
        "Other Deductions",
        "Total Deductions",
        "Net Salary",
        "Payment Status",
        "Payment Date",
    ]

    worksheet.append(headers)

    # Style the heading row.
    for cell in worksheet[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    # Add payroll data.
    for payroll, employee in records:
        worksheet.append([
            payroll.id,
            employee.employee_code,
            employee.full_name,
            employee.department,
            employee.designation,
            payroll.month,
            payroll.year,
            float(payroll.basic_salary or 0),
            float(payroll.hra or 0),
            float(payroll.da or 0),
            float(payroll.other_allowances or 0),
            float(payroll.gross_salary or 0),
            float(payroll.pf_deduction or 0),
            float(payroll.tax_deduction or 0),
            float(payroll.other_deductions or 0),
            float(payroll.total_deductions or 0),
            float(payroll.net_salary or 0),
            payroll.status,
            (
                payroll.payment_date.isoformat()
                if payroll.payment_date
                else ""
            ),
        ])

    # Format salary columns as currency amounts.
    salary_columns = range(8, 18)

    for row in worksheet.iter_rows(
        min_row=2,
        min_col=8,
        max_col=17,
    ):
        for cell in row:
            cell.number_format = '#,##0.00'

    # Automatically adjust column widths.
    for column_cells in worksheet.columns:
        column_letter = get_column_letter(
            column_cells[0].column
        )

        max_length = max(
            len(str(cell.value or ""))
            for cell in column_cells
        )

        worksheet.column_dimensions[
            column_letter
        ].width = min(max_length + 3, 30)

    # Freeze the header row.
    worksheet.freeze_panes = "A2"

    # Add Excel filters.
    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    # Prepare the workbook in memory.
    excel_buffer = io.BytesIO()
    workbook.save(excel_buffer)
    excel_buffer.seek(0)

    # Build a filename.
    filename = "payroll_report"

    if payroll_month is not None:
        filename += f"_month_{payroll_month:02d}"

    if payroll_year is not None:
        filename += f"_{payroll_year}"

    filename += ".xlsx"

    # Return the downloadable Excel file.
    return StreamingResponse(
        iter([excel_buffer.getvalue()]),
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="{filename}"'
            )
        },
    )